"""
Extraction du texte brut d'un CV uploadé (PDF ou DOCX).
Tout se fait en mémoire, rien n'est écrit sur disque.
"""
import io
import zipfile

import pdfplumber
from docx import Document

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 Mo


class ExtractionError(Exception):
    pass


def _sniff_kind(file_bytes: bytes) -> str | None:
    """
    Détermine le vrai type du fichier à partir de sa signature binaire
    (magic bytes), jamais à partir du Content-Type déclaré par le client
    — ce dernier est entièrement falsifiable par quiconque appelle l'API
    directement (sans passer par le formulaire web).
    """
    if file_bytes[:5] == b"%PDF-":
        return "pdf"
    if file_bytes[:4] == b"PK\x03\x04":
        # Un DOCX est un ZIP : on vérifie qu'il contient bien la structure
        # interne attendue, pas juste n'importe quelle archive ZIP renommée.
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                if "word/document.xml" in zf.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            return None
    return None


def extract_text(file_bytes: bytes, content_type: str = "") -> str:
    if len(file_bytes) > MAX_FILE_SIZE:
        raise ExtractionError("Fichier trop volumineux (5 Mo max).")

    kind = _sniff_kind(file_bytes)
    if kind is None:
        raise ExtractionError(
            "Format non supporté ou fichier invalide/corrompu. Utilisez un PDF ou un DOCX."
        )

    if kind == "pdf":
        return _extract_pdf(file_bytes)
    return _extract_docx(file_bytes)


def _extract_pdf(file_bytes: bytes) -> str:
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    text = "\n".join(text_parts).strip()
    if not text:
        raise ExtractionError(
            "Impossible d'extraire le texte du PDF (probablement un scan/image)."
        )
    return text


def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    text = "\n".join(parts).strip()
    if not text:
        raise ExtractionError("Le document DOCX semble vide.")
    return text
